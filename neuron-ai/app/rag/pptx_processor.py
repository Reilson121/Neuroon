from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


class PPTXProcessor:
    """
    Responsável pelo processamento inicial de arquivos PPTX.

    Funções atuais:
    - carregar apresentações;
    - extrair texto dos slides;
    - identificar elementos visuais;
    - detectar slides candidatos à análise multimodal.
    """

    def carregar(self, caminho_pptx):
        """
        Carrega um arquivo PPTX.
        """

        caminho = Path(caminho_pptx)

        if not caminho.exists():
            raise FileNotFoundError(
                f"PPTX não encontrado: {caminho}"
            )

        return Presentation(caminho)

    def extrair_texto_slide(self, slide):
        """
        Extrai todo o texto disponível diretamente
        nos objetos de um slide.
        """

        textos = []

        for shape in slide.shapes:

            if hasattr(shape, "text"):

                texto = shape.text.strip()

                if texto:
                    textos.append(texto)

        return "\n".join(textos)

    def analisar_slide(self, slide):
        """
        Analisa a estrutura básica do slide.

        Retorna:
        - texto extraído;
        - quantidade de caracteres;
        - quantidade de imagens;
        - quantidade de gráficos;
        - quantidade de outros elementos.
        """

        texto = self.extrair_texto_slide(slide)

        imagens = 0
        graficos = 0
        outros = 0

        for shape in slide.shapes:

            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                imagens += 1

            elif shape.shape_type == MSO_SHAPE_TYPE.CHART:
                graficos += 1

            else:
                outros += 1

        return {
            "texto": texto,
            "caracteres": len(texto),
            "imagens": imagens,
            "graficos": graficos,
            "outros_elementos": outros
        }

    def precisa_multimodal(self, slide):
        """
        Identifica slides candidatos à análise multimodal.

        Importante:
        Este método NÃO chama o modelo de IA.
        Ele apenas classifica o slide utilizando
        características estruturais do PowerPoint.
        """

        analise = self.analisar_slide(slide)

        caracteres = analise["caracteres"]
        imagens = analise["imagens"]
        graficos = analise["graficos"]
        outros = analise["outros_elementos"]

        motivos = []

        # Slide com pouco texto e pelo menos uma imagem.
        # A imagem pode conter informação importante
        # que não apareceu na extração textual.
        if caracteres < 180 and imagens > 0:
            motivos.append(
                "pouco_texto_com_imagem"
            )

        # Muitos objetos podem representar um grafo
        # construído diretamente no PowerPoint:
        # círculos, linhas, setas, caixas etc.
        if outros >= 10:
            motivos.append(
                "muitos_elementos_visuais"
            )

        # Gráficos nativos do PowerPoint também
        # merecem inspeção visual.
        if graficos > 0:
            motivos.append(
                "possui_grafico"
            )

        # Slide com pouquíssimo texto.
        if caracteres < 60:
            motivos.append(
                "texto_muito_curto"
            )

        return {
            **analise,
            "precisa_multimodal": len(motivos) > 0,
            "motivos": motivos
        }