from pathlib import Path
from pptx import Presentation


CAMINHO_PPTX = Path(
    "materiais_academicos/"
    "periodo_04/"
    "teoria_dos_grafos/"
    "aula_02_conceitos_basicos/"
    "Aula 02 - TG - Conceitos Básicos.pptx"
)

if not CAMINHO_PPTX.exists():
    raise FileNotFoundError(
        f"PPTX não encontrado: {CAMINHO_PPTX}"
    )


apresentacao = Presentation(CAMINHO_PPTX)

print(f"Total de slides: {len(apresentacao.slides)}")


for numero, slide in enumerate(apresentacao.slides, start=1):

    textos = []

    for shape in slide.shapes:

        if hasattr(shape, "text"):

            texto = shape.text.strip()

            if texto:
                textos.append(texto)

    conteudo = "\n".join(textos)

    print("\n" + "=" * 70)
    print(f"SLIDE {numero}")
    print("=" * 70)

    if conteudo:
        print(conteudo[:3000])
    else:
        print("[Nenhum texto extraído]")

    # Só os 5 primeiros no teste inicial
    if numero == 5:
        break